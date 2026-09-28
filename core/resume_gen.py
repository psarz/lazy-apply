
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document
import pathlib

def make_pdf(profile, tailored, out_path):
    doc = SimpleDocTemplate(out_path, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph(f"<b>{profile.get('full_name','')}</b>", styles['Title']))
    story.append(Paragraph(f"{profile.get('email','')} | {profile.get('phone','')} | {profile.get('location','')} | {profile.get('linkedin','')}", styles['Normal']))
    story.append(Spacer(1,12))
    story.append(Paragraph("<b>SUMMARY</b>", styles['Heading2']))
    story.append(Paragraph(tailored.get('summary',''), styles['Normal']))
    story.append(Spacer(1,12))
    story.append(Paragraph("<b>SKILLS</b>", styles['Heading2']))
    story.append(Paragraph(", ".join(tailored.get('skills_ranked', profile.get('skills',[]))), styles['Normal']))
    story.append(Spacer(1,12))
    story.append(Paragraph("<b>EXPERIENCE</b>", styles['Heading2']))
    for exp in tailored.get('experience', profile.get('experience',[])):
        story.append(Paragraph(f"<b>{exp.get('title','')} - {exp.get('company','')} ({exp.get('start','')}-{exp.get('end','')})</b>", styles['Heading3']))
        for b in exp.get('bullets',[]):
            story.append(Paragraph(f"• {b}", styles['Normal']))
    doc.build(story)
    return out_path

def make_docx(profile, tailored, out_path):
    doc = Document()
    doc.add_heading(profile.get('full_name',''), 0)
    doc.add_paragraph(f"{profile.get('email','')} | {profile.get('phone','')} | {profile.get('linkedin','')}")
    doc.add_heading('Summary', 2)
    doc.add_paragraph(tailored.get('summary',''))
    doc.add_heading('Skills', 2)
    doc.add_paragraph(", ".join(tailored.get('skills_ranked',[])))
    doc.add_heading('Experience', 2)
    for exp in tailored.get('experience',[]):
        doc.add_paragraph(f"{exp.get('title','')} - {exp.get('company','')}", style='Heading 3')
        for b in exp.get('bullets',[]):
            doc.add_paragraph(b, style='List Bullet')
    doc.save(out_path)
    return out_path
